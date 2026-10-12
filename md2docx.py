"""Convert a lecture summary .md (with {{A:}} {{D:}} {{T:}} markers) into a .docx
with coloured run shading. Usage: python md2docx.py in.md out.docx
Self-check: python md2docx.py --selftest
"""
import re
import sys

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt

FILL = {"A": "FFFF00", "D": "ADD8E6", "T": "D8BFD8"}  # yellow / light blue / light purple
TOKEN = re.compile(r"(\{\{[ADT]:[^{}]*\}\}|\*\*[^*]+\*\*|\*[^*]+\*)")


def shade(run, fill):
    rPr = run._element.get_or_add_rPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    rPr.append(shd)


def add_runs(p, text, bold=False, italic=False):
    for part in TOKEN.split(text):
        if not part:
            continue
        m = re.match(r"^\{\{([ADT]):([\s\S]*)\}\}$", part)
        if m:
            r = p.add_run(m.group(2))
            r.bold, r.italic = bold, italic
            shade(r, FILL[m.group(1)])
        elif part.startswith("**") and part.endswith("**"):
            add_runs(p, part[2:-2], True, italic)  # recurse: a marker can sit inside the bold
        elif len(part) > 2 and part.startswith("*") and part.endswith("*"):
            add_runs(p, part[1:-1], bold, True)
        else:
            r = p.add_run(part)
            r.bold, r.italic = bold, italic


def add_heading(doc, text, level):
    p = doc.add_paragraph(style="Heading %d" % level)
    add_runs(p, text)


def convert(src, dst):
    doc = Document()
    doc.styles["Normal"].font.size = Pt(11)
    for level, size in ((1, 18), (2, 16), (3, 14)):
        doc.styles["Heading %d" % level].font.size = Pt(size)
    with open(src, encoding="utf-8") as fh:
        lines = fh.read().splitlines()
    for raw in lines:
        line = raw.rstrip()
        if not line.strip():
            continue
        if line.startswith("### "):
            add_heading(doc, line[4:].strip(), 3)
        elif line.startswith("## "):
            add_heading(doc, line[3:].strip(), 2)
        elif line.startswith("# "):
            add_heading(doc, line[2:].strip(), 1)
        elif re.match(r"^\s*[-*] ", line):
            indent = (len(line) - len(line.lstrip())) // 2
            style = "List Bullet" if indent < 1 else "List Bullet %d" % min(indent + 1, 3)
            p = doc.add_paragraph(style=style)
            add_runs(p, re.sub(r"^\s*[-*] ", "", line))
        elif re.match(r"^\s*\d+\. ", line):
            p = doc.add_paragraph(style="List Number")
            add_runs(p, re.sub(r"^\s*\d+\. ", "", line))
        else:
            p = doc.add_paragraph()
            add_runs(p, line.strip())
    doc.save(dst)


def _selftest():
    import os
    import tempfile
    import zipfile

    src = tempfile.mktemp(suffix=".md")
    dst = tempfile.mktemp(suffix=".docx")
    with open(src, "w", encoding="utf-8") as fh:
        fh.write(
            "# Title\n\n## 1. Section\n"
            "- plain bullet\n"
            "- {{A:Smith (2001)}} and {{T:Resource Mobilization}} mean "
            "{{D:the pooling of resources by a social movement organisation}}.\n"
            "\n### Sub\nA **bold** line.\n"
            "- **{{T:Bold Theory}}** sits inside bold.\n"
            "- *Pride and Prejudice* is italic, and *{{A:Smith (2001)}}* is both.\n"
        )
    convert(src, dst)
    xml = zipfile.ZipFile(dst).read("word/document.xml").decode("utf-8")
    for fill in FILL.values():
        assert fill in xml, "missing shading %s" % fill
    assert "{{" not in xml, "marker leaked into output"
    assert "bold" in xml, "body text missing"
    assert "*" not in xml, "stray asterisk leaked into output"
    assert "Pride and Prejudice" in xml, "italic text missing"
    os.remove(src)
    os.remove(dst)
    print("selftest OK")


if __name__ == "__main__":
    if len(sys.argv) == 2 and sys.argv[1] == "--selftest":
        _selftest()
    else:
        convert(sys.argv[1], sys.argv[2])
